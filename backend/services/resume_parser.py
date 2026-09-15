import fitz
from docx import Document
from io import BytesIO


def extract_text(filename: str, file_content: bytes):

    if filename.lower().endswith(".pdf"):

        pdf = fitz.open(
            stream=file_content,
            filetype="pdf"
        )

        text = ""

        for page in pdf:
            text += page.get_text()

        pdf.close()

        return text


    elif filename.lower().endswith(".docx"):

        document = Document(
            BytesIO(file_content)
        )

        text = ""


        # Read normal paragraphs

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text += paragraph.text + "\n"


        # Read tables

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text.strip():
                        row_text.append(cell.text.strip())

                if row_text:
                    text += " | ".join(row_text) + "\n"


        return text


    else:

        raise ValueError(
            "Unsupported file type"
        )